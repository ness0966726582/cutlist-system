"""匯入既有 Excel 料單 (拆單系統 2026年6.0版 格式) → 專案 / 櫃體 / 板材。

解析「料單」分頁：表頭 (C1 訂單編號, C2 案場, C3 業務/設計師, K2 訂貨日期, K3 交貨日期, L3 拉/KD, K4 送貨地址)
與每個櫃體區塊 (A 欄 = 櫃體名稱 的列開始：名稱 C, 櫃號 C+1, 高/深/寬 C+2..C+4, 批次 L+3, 明細列 +7 起)。

規則：
- 明細列只認「數量 > 0」的列 (與 Excel 片數 SUMIF 同邏輯)，範本預填但沒填數量的列不算。
- 板材序號依列的順序重新編 (櫃號×1000+序)，不信任 C 欄 (可能被手改、重複或錯位)。
- 同櫃號的多個區塊 (匯出的「(續)」區塊或手動拆的) 合併成一個櫃體，名稱取第一個區塊，去掉尾端「(續)」。
- 有名稱或尺寸但沒有板材的櫃體仍保留。
"""
import collections
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from . import rules

JOINT_TYPES = {"L", "KD", "拉", "拉米諾"}


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def parse_workbook(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "料單" not in wb.sheetnames:
        raise ValueError("找不到「料單」分頁")
    ws = wb["料單"]
    joint = rules.norm_text(ws["L3"].value)
    header = {
        "order_no": rules.norm_text(ws["C1"].value) or None,
        "site_name": rules.norm_text(ws["C2"].value) or "未命名案場",
        "designer": rules.norm_text(ws["C3"].value) or None,
        "order_date": _date(ws["K2"].value),
        "delivery_date": _date(ws["K3"].value),
        "joint_type": joint if joint in JOINT_TYPES else None,
        "delivery_address": rules.norm_text(ws["K4"].value) or rules.norm_text(ws["L4"].value) or None,
    }
    cabinets: dict[int, dict] = {}
    order: list[int] = []
    warnings: list[str] = []
    thick_counter = collections.Counter()
    r = 1
    max_row = ws.max_row
    while r <= max_row:
        if ws.cell(r, 1).value == "櫃體名稱":
            name = re.sub(r"\(續\)$", "", rules.norm_text(ws.cell(r, 3).value)).strip()
            no = _num(ws.cell(r + 1, 3).value)
            dims = [_num(ws.cell(r + k, 3).value) for k in (2, 3, 4)]
            parts = []
            rr = r + 7
            while rr <= max_row and ws.cell(rr, 1).value != "櫃體名稱":
                nm = rules.norm_text(ws.cell(rr, 5).value)
                qty = _num(ws.cell(rr, 11).value)
                if nm and qty and qty > 0:
                    cat = rules.norm_text(ws.cell(rr, 4).value) or None
                    thick = _num(ws.cell(rr, 10).value)
                    parts.append({
                        "excel_part_no": _num(ws.cell(rr, 3).value),
                        "category": cat,
                        "name": nm,
                        "length": _num(ws.cell(rr, 6).value),
                        "edge_l": rules.norm_edge(ws.cell(rr, 7).value),
                        "width": _num(ws.cell(rr, 8).value),
                        "edge_w": rules.norm_edge(ws.cell(rr, 9).value),
                        "thickness": thick,
                        "qty": int(qty),
                        "color_code": rules.norm_text(ws.cell(rr, 12).value) or None,
                        "mark_p": rules.norm_text(ws.cell(rr, 16).value),
                        "mark_q": rules.norm_text(ws.cell(rr, 17).value),
                        "mark_r": rules.norm_text(ws.cell(rr, 18).value),
                        "remark": rules.norm_text(ws.cell(rr, 19).value) or None,
                        "reminder": rules.norm_text(ws.cell(rr, 20).value),
                    })
                    if thick and not (cat or "").endswith("五金") and (parts[-1]["length"] or parts[-1]["width"]):
                        thick_counter[thick] += 1
                elif nm and (ws.cell(rr, 6).value not in (None, "") or ws.cell(rr, 8).value not in (None, "")) and (name or any(d is not None for d in dims)):
                    if _num(ws.cell(rr, 11).value) is not None:
                        warnings.append(f"第 {rr} 列「{nm}」數量為 0，未匯入")
                rr += 1
            if name or any(d is not None for d in dims) or parts:
                key = int(no) if no else (max(order) + 1 if order else 1)
                if key in cabinets:
                    cab = cabinets[key]
                    cab["parts"].extend(parts)
                    if not cab["name"] and name:
                        cab["name"] = name
                else:
                    cabinets[key] = {"name": name or f"櫃 {key}", "cabinet_no": key, "height": dims[0], "depth": dims[1],
                                     "width": dims[2], "batch": int(_num(ws.cell(r + 3, 12).value) or 1), "parts": parts}
                    order.append(key)
            r = rr
        else:
            r += 1
    # 背板 4mm 不算；取非五金板材最常見的厚度 (17.5 / 18)
    body = {k: v for k, v in thick_counter.items() if k not in (4, 4.0)}
    header["board_thickness"] = (max(body, key=body.get) if body else 18)
    return {"header": header, "cabinets": [cabinets[k] for k in order], "warnings": warnings}
