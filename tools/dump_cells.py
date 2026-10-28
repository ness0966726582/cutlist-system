import sys
from openpyxl import load_workbook
path, out = sys.argv[1], sys.argv[2]
wb = load_workbook(path, data_only=False)
lines = []
def cell(sheet, ref):
    v = wb[sheet][ref].value
    if hasattr(v, "text"): v = "{ARR}" + str(v.text)
    lines.append(f"### {sheet}!{ref}\n{v}\n")
for ref in ["G14", "I14", "P14", "Q14", "R14", "T14", "D7", "C1"]:
    cell("料單", ref)
ws = wb["料單"]
for r in list(range(40, 53)) + list(range(82, 92)):
    cells = []
    for c in range(1, 24):
        v = ws.cell(r, c).value
        if v is None: continue
        if hasattr(v, "text"): v = "{ARR}" + str(v.text)
        v = " ".join(str(v).split())
        cells.append(f"{ws.cell(r,c).coordinate}={v[:60]}")
    if cells: lines.append(f"R{r}: " + " | ".join(cells))
for name, n in [("單價表", 6), ("色號", 12), ("客戶資料", 10), ("標籤列印版", 3)]:
    ws = wb[name]
    lines.append(f"\n## {name}")
    for r in range(1, n + 1):
        cells = []
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(r, c).value
            if v is None: continue
            if hasattr(v, "text"): v = "{ARR}" + str(v.text)
            v = " ".join(str(v).split())
            cells.append(f"{ws.cell(r,c).coordinate}={v[:70]}")
        if cells: lines.append(f"R{r}: " + " | ".join(cells))
# counts of master data rows
for name, col in [("色號", 1), ("單價表", 2), ("客戶資料", 2), ("五金配件", 4), ("資料庫", 3)]:
    ws = wb[name]
    n = sum(1 for r in range(1, ws.max_row + 1) if ws.cell(r, col).value not in (None, ""))
    lines.append(f"rows with data in {name} col{col}: {n}")
open(out, "w", encoding="utf-8").write("\n".join(lines))
