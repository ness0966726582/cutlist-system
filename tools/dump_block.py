import sys
from openpyxl import load_workbook
path, sheet, r1, r2, out = sys.argv[1], sys.argv[2], int(sys.argv[3]), int(sys.argv[4]), sys.argv[5]
wb = load_workbook(path, data_only=False)
ws = wb[sheet]
lines = []
for r in range(r1, r2 + 1):
    cells = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v is None:
            continue
        if hasattr(v, "text"):
            v = "{ARR}" + str(v.text)
        v = " ".join(str(v).split())
        cells.append(f"{ws.cell(r,c).coordinate}={v[:160]}")
    if cells:
        lines.append(f"R{r}: " + " | ".join(cells))
open(out, "w", encoding="utf-8").write("\n".join(lines))
print(len(lines))
