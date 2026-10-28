"""Dump filled-in values (data_only) of a 料單 workbook: header, cabinets, non-empty part rows."""
import sys
from openpyxl import load_workbook

path, out = sys.argv[1], sys.argv[2]
wb = load_workbook(path, data_only=True)
lines = []
print("sheets:", wb.sheetnames)
ws = wb["料單"]
lines.append("== 料單 header")
for r in range(1, 6):
    cells = [f"{ws.cell(r,c).coordinate}={ws.cell(r,c).value}" for c in range(1, 22) if ws.cell(r, c).value not in (None, "", " ")]
    lines.append(" | ".join(cells))

# scan cabinet blocks: header row where A == '櫃體名稱'
r = 1
blocks = []
while r <= ws.max_row:
    if ws.cell(r, 1).value == "櫃體名稱":
        blocks.append(r)
    r += 1
lines.append(f"blocks: {len(blocks)}")
for b in blocks:
    name = ws.cell(b, 3).value
    no = ws.cell(b + 1, 3).value
    h, d, w = ws.cell(b + 2, 3).value, ws.cell(b + 3, 3).value, ws.cell(b + 4, 3).value
    pieces, cai, batch = ws.cell(b + 1, 12).value, ws.cell(b + 2, 12).value, ws.cell(b + 3, 12).value
    note = ws.cell(b, 4).value
    rows = []
    rr = b + 7
    while rr <= ws.max_row and ws.cell(rr, 1).value != "櫃體名稱" and rr < b + 60:
        e = ws.cell(rr, 5).value
        if e not in (None, ""):
            rows.append([ws.cell(rr, c).value for c in range(3, 22)])
        rr += 1
    if name in (None, "") and not rows:
        continue
    lines.append(f"\n-- block@R{b} 櫃體={name} 櫃號={no} 高={h} 深={d} 寬={w} 片數={pieces} 才數={cai} 批次={batch} 備註={note}")
    for row in rows:
        lines.append("   " + " | ".join("" if v is None else str(v).replace("\n", " ") for v in row))

# other sheets with user content: 零件單, 計價清價 (2) header
for name in ["零件單"]:
    ws2 = wb[name]
    lines.append(f"\n== {name}")
    for r in range(1, ws2.max_row + 1):
        cells = [f"{ws2.cell(r,c).coordinate}={ws2.cell(r,c).value}" for c in range(1, ws2.max_column + 1) if ws2.cell(r, c).value not in (None, "")]
        if cells:
            lines.append(" | ".join(cells))
ws3 = wb["計價清價 (2)"]
lines.append("\n== 計價清價 (2) header")
for r in range(1, 12):
    cells = [f"{ws3.cell(r,c).coordinate}={ws3.cell(r,c).value}" for c in range(1, 19) if ws3.cell(r, c).value not in (None, "")]
    if cells:
        lines.append(" | ".join(cells))
open(out, "w", encoding="utf-8").write("\n".join(lines))
print("lines", len(lines))
