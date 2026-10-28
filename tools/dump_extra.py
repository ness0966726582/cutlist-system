import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
L = []
def dump(sheet, r1, r2, c1, c2):
    ws = wb[sheet]
    L.append(f"\n## {sheet} rows {r1}-{r2} cols {c1}-{c2}")
    for r in range(r1, r2 + 1):
        cells = []
        for c in range(c1, c2 + 1):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                cells.append(f"{ws.cell(r,c).coordinate}={str(v).replace(chr(10),' ')[:40]}")
        if cells:
            L.append(" | ".join(cells))
dump("資料庫", 1, 70, 29, 40)      # AC..AN
dump("牌價表(對公司內部)", 1, 40, 1, 33)
dump("對外賣價表(無連結)", 1, 12, 1, 6)
dump("20230916成本分析", 5, 50, 3, 25)
dump("報價單", 1, 39, 1, 8)
dump("客戶對帳單", 1, 33, 1, 9)
dump("零件單", 1, 10, 1, 4)
dump("色號1", 1, 3, 1, 18)
open(sys.argv[2], "w", encoding="utf-8").write("\n".join(L))
print(len(L))
