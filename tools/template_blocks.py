import sys, time
from openpyxl import load_workbook

path = sys.argv[1]
t = time.time()
wb = load_workbook(path)
print("load s", round(time.time() - t, 1))
ws = wb["料單"]
starts = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "櫃體名稱"]
print("blocks", len(starts))
caps = []
for i, b in enumerate(starts):
    end = starts[i + 1] - 1 if i + 1 < len(starts) else ws.max_row
    # part rows: from b+7 while C has formula/value (numeric-ish) ; find last row with C formula before end
    last = b + 7
    for r in range(b + 7, end + 1):
        if ws.cell(r, 3).value not in (None, ""):
            last = r
    caps.append((b, b + 7, last, last - (b + 7) + 1))
print(caps[:6], "...", caps[-3:])
print("distinct caps", sorted({c[3] for c in caps}))
# header cells near 修改日期
for ref in ["N2", "O2", "P2", "Q2", "N3", "O3", "N4", "O4", "L3", "L4", "K2", "K3", "G1", "C1"]:
    print(ref, repr(ws[ref].value))
# row 44 / block gap content
b = starts[1]
for r in range(b - 2, b + 1):
    print(r, [(ws.cell(r, c).coordinate, str(ws.cell(r, c).value)[:30]) for c in range(1, 22) if ws.cell(r, c).value is not None])
t = time.time()
wb.save(sys.argv[2])
print("save s", round(time.time() - t, 1))
