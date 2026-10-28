"""Regression checks for the bug-review fixes (uses throwaway project, deleted at end)."""
import json, sys, urllib.request, uuid

B = "http://localhost:8100/api"
out = []
log = lambda *a: out.append(" ".join(str(x) for x in a))


def call(method, path, body=None):
    req = urllib.request.Request(B + path, method=method, headers={"Content-Type": "application/json"},
                                 data=json.dumps(body).encode() if body is not None else None)
    try:
        with urllib.request.urlopen(req) as r:
            t = r.read().decode(); return r.status, (json.loads(t) if t else None)
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


def upload(path, filepath, fname):
    b = uuid.uuid4().hex
    data = open(filepath, "rb").read()
    body = (f'--{b}\r\nContent-Disposition: form-data; name="file"; filename="{fname}"\r\nContent-Type: application/octet-stream\r\n\r\n').encode() + data + f"\r\n--{b}--\r\n".encode()
    req = urllib.request.Request(B + path, method="POST", data=body, headers={"Content-Type": f"multipart/form-data; boundary={b}"})
    try:
        with urllib.request.urlopen(req) as r:
            return r.status, json.loads(r.read().decode())
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


s, p = call("POST", "/projects", {"site_name": "回歸測試", "board_thickness": 18}); pid = p["id"]
s, r = call("POST", "/projects", {"site_name": "  "}); log("#28 blank site_name ->", s)
s, cb = call("POST", f"/projects/{pid}/cabinets", {"name": "高櫃", "height": 2000, "depth": 600, "width": 1200, "color_code": "1774"})
log("#15 default parts colors:", {x["name"]: x["color_code"] for x in cb["parts"][:5]})
s, r = call("POST", f"/projects/{pid}/cabinets", {"name": "重複", "cabinet_no": 1}); log("#6 duplicate cabinet_no ->", s, str(r)[:60])
s, r = call("POST", f"/projects/{pid}/cabinets", {"name": "零", "cabinet_no": 0}); log("#18 cabinet_no 0 ->", s)
# #1 inference
s, r = call("POST", f"/cabinets/{cb['id']}/parts", {"name": "頂板", "length": 500, "width": 400}); log("#1 頂板 category:", r["category"], "#13 thickness/color defaults:", r["thickness"], r["color_code"])
s, r = call("POST", f"/cabinets/{cb['id']}/parts", {"name": "左側板"}); log("#1 左側板 category:", r["category"])
s, hw = call("POST", f"/cabinets/{cb['id']}/parts", {"name": "活隔粒", "qty": 8}); log("#12 hardware thickness/color:", hw["category"], hw["thickness"], hw["color_code"])
# #11 rename re-infer
body = {k: r[k] for k in ("category", "name", "length", "width", "thickness", "qty", "color_code", "remark", "status_code")}
body["name"] = "拉米諾"; s, r2 = call("PUT", f"/parts/{r['id']}", body); log("#11 rename 左側板->拉米諾 category:", r2["category"])
# #14 invalid status
body["status_code"] = "ZZ"; s, r3 = call("PUT", f"/parts/{r['id']}", body); log("#14 invalid status ->", s)
# #7 empty-name draft
s, d = call("POST", f"/cabinets/{cb['id']}/parts", {"name": ""}); log("#7 draft created:", d["id"], repr(d["name"]))
s, lb = call("GET", f"/projects/{pid}/labels"); log("#7 labels exclude draft:", all(x["name"] for x in lb), len(lb))
s, lst = call("GET", "/projects"); me = next(x for x in lst if x["id"] == pid); log("#7 piece_count:", me["piece_count"])
# #2 serial stability
s, se = call("GET", f"/projects/{pid}/serials"); log("#2 serials:", [(x["serial_no"], x["name"], x["piece"]) for x in se[:4]], "n=", len(se))
s, r = call("POST", f"/projects/{pid}/scans", {"serial_no": 3, "stage": "cut"}); log("scan 3:", s, r)
first = cb["parts"][0]; call("DELETE", f"/parts/{first['id']}")
s, se2 = call("GET", f"/projects/{pid}/serials"); three = next(x for x in se2 if x["serial_no"] == 3)
log("#2 after deleting serial-1 part, serial 3 is still:", se[2]["name"], "->", three["name"], "cut=", three["cut"], "serial 1 gone:", all(x["serial_no"] != 1 for x in se2))
s, r = call("POST", f"/projects/{pid}/scans", {"serial_no": 99999, "stage": "cut"}); log("#14 bad serial ->", s)
s, r = call("POST", "/projects/999999/scans", {"serial_no": 1, "stage": "cut"}); log("#14 bad project ->", s)
s, r = call("DELETE", f"/projects/{pid}/scans/3/cut"); log("unscan ->", s)
# #16 totals agree
s, sm = call("GET", f"/projects/{pid}/summary"); s, qt = call("GET", f"/projects/{pid}/quote")
log("#16 summary total", sm["total_amount"], "quote subtotal", qt["subtotal"])
# #32 duplicate clears drawing label; #19 sort
s, r = call("PUT", f"/cabinets/{cb['id']}", {"name": "水槽櫃", "height": 2000, "depth": 600, "width": 1200, "drawing_label": "①", "generate_default_parts": False})
log("#20 note auto on rename to 水槽櫃:", r["note"])
s, dup = call("POST", f"/cabinets/{cb['id']}/duplicate"); log("#32 dup label:", dup["drawing_label"], "no:", dup["cabinet_no"], "sort:", dup["sort"])
# #31 material hardware key
s, mt = call("GET", f"/projects/{pid}/material"); log("#31 hardware:", mt["hardware"])
# #29 catalog duplicate
s, r = call("POST", "/master/parts-catalog", {"category": "桶身", "name": "頂板"}); log("#29 dup catalog ->", s)
s, r = call("POST", "/master/parts-catalog", {"category": "不存在", "name": "x"}); log("#29 bad category ->", s)
# #3/#4/#5 import with duplicate cabinet numbers & (續)
import io
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
ws = wb["料單"]
# block 2 (R45) gets same 櫃號 as block 1, name with (續); block1 parts rows 14.. ; give block2 row 52 a part
ws["C45"] = "測試櫃(續)"; ws["C46"] = 1
ws["C7"] = "測試櫃"; ws["C8"] = 1; ws["C9"] = 2000; ws["C10"] = 600; ws["C11"] = 1200
for r_, nm in ((14, "頂板"), (15, "底板")):
    ws.cell(r_, 5).value = nm; ws.cell(r_, 6).value = 1000; ws.cell(r_, 8).value = 500; ws.cell(r_, 11).value = 1; ws.cell(r_, 3).value = "1004"
ws.cell(16, 5).value = "左側板"; ws.cell(16, 11).value = None   # qty blank -> not a part (#8)
ws.cell(52, 5).value = "層板"; ws.cell(52, 6).value = 900; ws.cell(52, 8).value = 300; ws.cell(52, 11).value = 2
ws["L3"] = "拉、KD"
buf = io.BytesIO(); wb.save(buf)
open(sys.argv[2], "wb").write(buf.getvalue())
s, imp = upload("/projects/import-xlsx", sys.argv[2], "merge.xlsx"); log("#3/#4 import ->", s, imp)
s, pi = call("GET", f"/projects/{imp['project_id']}")
log("#3/#4 cabinets:", [(c["cabinet_no"], c["name"], [(x["part_no"], x["name"], x["qty"]) for x in c["parts"]]) for c in pi["cabinets"]], "#22 joint:", pi["joint_type"], "thick:", pi["board_thickness"])
call("DELETE", f"/projects/{imp['project_id']}")
call("DELETE", f"/projects/{pid}")
open(sys.argv[3], "w", encoding="utf-8").write("\n".join(out)); print("done")
