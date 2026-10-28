"""Download export of a project, inspect it, re-import it and compare totals."""
import json, sys, time, urllib.request, uuid
from openpyxl import load_workbook

B = "http://localhost:8100/api"
pid, out_xlsx = sys.argv[1], sys.argv[2]
t = time.time()
with urllib.request.urlopen(f"{B}/projects/{pid}/export-xlsx") as r:
    data = r.read(); cd = r.headers.get("content-disposition")
print("export bytes", len(data), "in", round(time.time() - t, 1), "s;", cd)
open(out_xlsx, "wb").write(data)

wb = load_workbook(out_xlsx)
ws = wb["料單"]
print("sheets", len(wb.sheetnames), "calc fullCalcOnLoad", wb.calculation.fullCalcOnLoad)
for ref in ("C1", "C2", "C3", "K2", "K3", "L3", "L4", "C4", "B5"):
    print(ref, repr(ws[ref].value))
for r in (7, 8, 9, 10, 11, 10 + 0):
    print(r, [(ws.cell(r, c).coordinate, str(ws.cell(r, c).value)[:25]) for c in (1, 3, 4, 12) if ws.cell(r, c).value is not None])
for r in range(14, 32):
    vals = [ws.cell(r, c).value for c in range(3, 21)]
    if any(v is not None for v in vals):
        print(r, [str(v)[:18] if v is not None else "" for v in vals])
print("U14", ws["U14"].value, "| A14", ws["A14"].value, "| L8", ws["L8"].value)
# block 3 (樓梯收納矮櫃): manual edges should be 'A'
b3 = 84
print("blk3", ws.cell(b3, 3).value, [(ws.cell(r, 5).value, ws.cell(r, 7).value, ws.cell(r, 9).value, ws.cell(r, 16).value, ws.cell(r, 17).value) for r in range(b3 + 7, b3 + 11)])
# untouched template block after last used
starts = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "櫃體名稱"]
used = [s for s in starts if ws.cell(s, 3).value]
print("used blocks", len(used), "next block name cell", ws.cell(starts[len(used)], 3).value, "E first row", ws.cell(starts[len(used)] + 7, 5).value)
print("標籤列印版 B2", wb["標籤列印版"]["B2"].value, "| 板材才數分析 O2", wb["板材才數分析."]["O2"].value)

# re-import exported file and compare totals
bnd = uuid.uuid4().hex
body = (f'--{bnd}\r\nContent-Disposition: form-data; name="file"; filename="export.xlsx"\r\nContent-Type: application/octet-stream\r\n\r\n').encode() + data + f"\r\n--{bnd}--\r\n".encode()
r2 = json.loads(urllib.request.urlopen(urllib.request.Request(f"{B}/projects/import-xlsx", data=body, headers={"Content-Type": f"multipart/form-data; boundary={bnd}"})).read())
print("reimport", r2)
s1 = json.loads(urllib.request.urlopen(f"{B}/projects/{pid}/summary").read())
s2 = json.loads(urllib.request.urlopen(f"{B}/projects/{r2['project_id']}/summary").read())
print("orig  pieces", s1["total_pieces"], "cai", s1["total_cai"], "amount", s1["total_amount"])
print("round pieces", s2["total_pieces"], "cai", s2["total_cai"], "amount", s2["total_amount"])
p1 = json.loads(urllib.request.urlopen(f"{B}/projects/{pid}").read())
p2 = json.loads(urllib.request.urlopen(f"{B}/projects/{r2['project_id']}").read())
diff = 0
for c1, c2 in zip(p1["cabinets"], p2["cabinets"]):
    for a, b in zip(c1["parts"], c2["parts"]):
        for k in ("name", "length", "width", "thickness", "qty", "color_code", "edge_l", "edge_w", "mark_p", "mark_q", "mark_r", "remark", "reminder"):
            if (a[k] or "") != (b[k] or "") and not (isinstance(a[k], (int, float)) and isinstance(b[k], (int, float)) and float(a[k]) == float(b[k])):
                diff += 1
                if diff < 8:
                    print("  diff", c1["cabinet_no"], a["part_no"], k, repr(a[k]), "->", repr(b[k]))
print("field diffs after roundtrip:", diff)
urllib.request.urlopen(urllib.request.Request(f"{B}/projects/{r2['project_id']}", method="DELETE")).read()
print("cleaned up roundtrip project")
