"""Aggregate how the real case workbook deviates from the blank template (for UI design)."""
import sys, collections
from openpyxl import load_workbook

case, blank, out = sys.argv[1], sys.argv[2], sys.argv[3]
wc = load_workbook(case, data_only=True)
wbk = load_workbook(blank, data_only=True)
L = []

# catalog from blank 拆單用
cat = collections.defaultdict(set)
ws = wbk["拆單用"]
for c in range(1, ws.max_column + 1):
    k = ws.cell(2, c).value
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, c).value
        if v:
            cat[k].add(" ".join(str(v).split()))

ws = wc["料單"]
names_not_in_cat = collections.Counter()
cats_used = collections.Counter()
edges = collections.Counter()
P, Q, R, S, T = (collections.Counter() for _ in range(5))
thick = collections.Counter()
colors = collections.Counter()
qty_gt1_body = 0
cab_names = collections.Counter()
cab_count = 0
for r in range(1, ws.max_row + 1):
    if ws.cell(r, 1).value == "櫃體名稱" and ws.cell(r, 3).value:
        cab_count += 1
        cab_names[ws.cell(r, 3).value] += 1
    e = ws.cell(r, 5).value
    if not e or ws.cell(r, 3).value in (None, "專案") or not isinstance(ws.cell(r, 3).value, (int, float)):
        continue
    d = ws.cell(r, 4).value
    nm = " ".join(str(e).split())
    cats_used[d] += 1
    if d not in cat or nm not in cat[d]:
        names_not_in_cat[(d, nm)] += 1
    for col, ctr in ((7, edges), (9, edges)):
        v = ws.cell(r, col).value
        if v not in (None, ""):
            ctr[repr(v)] += 1
    for col, ctr in ((16, P), (17, Q), (18, R), (19, S), (20, T)):
        v = ws.cell(r, col).value
        if v not in (None, ""):
            ctr[" ".join(str(v).split())] += 1
    if ws.cell(r, 10).value not in (None, ""):
        thick[ws.cell(r, 10).value] += 1
    if ws.cell(r, 12).value not in (None, ""):
        colors[str(ws.cell(r, 12).value)] += 1
    if d == "桶身" and (ws.cell(r, 11).value or 0) > 1:
        qty_gt1_body += 1

L.append(f"cabinets used: {cab_count}; duplicate names: {[(k,v) for k,v in cab_names.items() if v>1]}")
L.append(f"categories used: {dict(cats_used)}")
L.append(f"names NOT in catalog ({len(names_not_in_cat)}): {sorted(names_not_in_cat.items(), key=lambda x:-x[1])}")
L.append(f"edge values: {dict(edges)}")
L.append(f"thickness: {dict(thick)}")
L.append(f"colors: {dict(colors)}")
L.append(f"桶身 rows with qty>1: {qty_gt1_body}")
L.append(f"P values: {dict(P)}")
L.append(f"Q values: {dict(Q)}")
L.append(f"R values: {dict(R)}")
L.append(f"S values: {dict(S)}")
L.append(f"T values: {dict(T)}")

# master data additions vs blank
for sheet, keycol in (("色號", 1), ("拆單用", None), ("標記", 1)):
    a, b = wbk[sheet], wc[sheet]
    if keycol:
        ka = {a.cell(r, keycol).value for r in range(1, a.max_row + 1)}
        added = [[b.cell(r, c).value for c in range(1, min(b.max_column, 7) + 1)] for r in range(1, b.max_row + 1) if b.cell(r, keycol).value not in ka and b.cell(r, keycol).value not in (None, "")]
        L.append(f"{sheet} added rows ({len(added)}): {added[:40]}")
    else:
        added = []
        for c in range(1, b.max_column + 1):
            k = b.cell(2, c).value
            for r in range(3, b.max_row + 1):
                v = b.cell(r, c).value
                if v and " ".join(str(v).split()) not in cat.get(k, set()):
                    added.append((k, v))
        L.append(f"拆單用 added ({len(added)}): {added}")
open(out, "w", encoding="utf-8").write("\n".join(L))
print("ok")
