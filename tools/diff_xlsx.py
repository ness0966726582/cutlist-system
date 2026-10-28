import sys, zipfile, re
from openpyxl import load_workbook

a, b = sys.argv[1], sys.argv[2]
wa = load_workbook(a, data_only=False)
wb = load_workbook(b, data_only=False)

print("== sheet lists")
sa, sb = wa.sheetnames, wb.sheetnames
print("only in A:", [s for s in sa if s not in sb])
print("only in B:", [s for s in sb if s not in sa])

print("== defined names")
da = {k: v.attr_text for k, v in wa.defined_names.items()}
db = {k: v.attr_text for k, v in wb.defined_names.items()}
for k in sorted(set(da) | set(db)):
    if da.get(k) != db.get(k):
        print(" ", k, "|", da.get(k), "->", db.get(k))
for k in sorted(da):
    print("  name:", k, "=", da[k])

print("== cell diffs (A -> B)")
total = 0
for name in sa:
    if name not in sb:
        continue
    x, y = wa[name], wb[name]
    rows = max(x.max_row, y.max_row)
    cols = max(x.max_column, y.max_column)
    diffs = []
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            va = x.cell(r, c).value
            vb = y.cell(r, c).value
            if hasattr(va, "text"): va = "{ARR}" + str(va.text)
            if hasattr(vb, "text"): vb = "{ARR}" + str(vb.text)
            if va != vb:
                diffs.append((x.cell(r, c).coordinate, va, vb))
    if diffs:
        total += len(diffs)
        print(f"-- {name}: {len(diffs)} diffs")
        for d in diffs[:40]:
            print(f"   {d[0]}: {str(d[1])[:120]!r} -> {str(d[2])[:120]!r}")
print("total diffs:", total)

print("== external links / connections (A)")
z = zipfile.ZipFile(a)
for n in z.namelist():
    if n.startswith("xl/externalLinks/_rels/"):
        print(n, re.findall(r'Target="([^"]+)"', z.read(n).decode("utf-8", "ignore")))
    if n == "xl/connections.xml":
        s = z.read(n).decode("utf-8", "ignore")
        print(n, re.findall(r'<connection [^>]*name="([^"]+)"', s))
        print(re.findall(r'connection="([^"]+)"', s)[:5])
        print(re.findall(r'command="([^"]+)"', s)[:5])
    if n.startswith("xl/tables/table") and n.endswith(".xml"):
        s = z.read(n).decode("utf-8", "ignore")
        print(n, re.findall(r'<table [^>]*displayName="([^"]+)"[^>]*ref="([^"]+)"', s), re.findall(r'<tableColumn [^>]*name="([^"]+)"', s)[:30])
